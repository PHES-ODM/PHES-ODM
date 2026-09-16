"""
Generate dictionary-tables/*.csv from the team's xlsx dictionary-authoring
workbook. The inverse of generate_workbook.py in this same folder (that one
goes CSV -> xlsx; this one goes xlsx -> CSV). Not part of the regular workflow
(dictionary-tables/*.csv is the source of truth going forward -- see AGENTS.md/
CLAUDE.md), but kept correct and reproducible for the occasional case an xlsx
snapshot needs re-importing, and as a record of exactly what that requires.

Reads with openpyxl directly, not pandas -- deliberately. Two real bugs this
session both trace back to pandas' Excel reader collapsing distinct states into
a single NaN, with no way to tell them apart afterward:
  1. Literal "NA"/"NULL" text (real data in some cells, e.g. Namibia's ISO code,
     or the maintainer's deliberate "not applicable" marker elsewhere -- see
     CLAUDE.md, "the dictionary CSVs ... are the data") gets silently
     re-interpreted as missing by pandas' default na_values list, even with
     dtype=str. openpyxl never does this -- a cell containing the text "NA"
     comes back as the string "NA", full stop.
  2. A formula cell whose LAST CACHED VALUE was itself an Excel error (#N/A,
     #REF!, etc. -- see the wideName column: two rows had a stale #N/A cached
     from an earlier broken formula-dependency state, long since fixed in a
     later save of the workbook) also reads back as NaN via pandas, identical
     to a cell that was simply never filled in. openpyxl's data_only=True
     surfaces the error as the literal string "#N/A" instead, which is exactly
     what makes it possible to tell "genuinely blank" and "errored formula"
     apart -- and this script does, loudly, rather than silently dropping the
     row the way the original export from this same class of workbook did
     (two ODM_wideNames.csv rows went missing this way and had to be manually
     recovered from a later xlsx snapshot -- see docs/dictionary-rules/
     parts-by-partType.md item 19 and the September 2026 session notes).
"""
import csv
import os
import re
import sys
import openpyxl
import warnings

warnings.filterwarnings("ignore")

TARGET_VERSION = "3.0.1"
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), "dictionary-tables")

DE_DW_TRIMMED = {
    "label_fr", "partDescription_fr", "partInstruction_fr",
    "nwss", "nwssProcess", "nwssNotes", "ena", "enaNotes",
    "norman", "normanNotes", "wSphere", "wSphereNotes", "ncbi", "phage",
    "version1Table", "version1Location", "version1Variable",
    "version1Category", "version1to2Changes",
}

WS_RE = re.compile(r"^[\s\xa0]+|[\s\xa0]+$")

# sheet name -> (output basename, primary key column name)
SHEETS = {
    "parts": ("ODM_parts", "partID"),
    "sets": ("ODM_sets", "setCompID"),
    "wideNames": ("ODM_wideNames", "wideName"),
    "languages": ("ODM_languages", "lang"),
    "translations": ("ODM_translations", "translationID"),
    "countries": ("ODM_countries", "isoCode"),
    "zones": ("ODM_zones", "isoCode"),
}

# translations.csv is deliberately NEVER written by this script (maintainer's call,
# September 2026): its content has diverged from anything derivable off an xlsx
# snapshot -- most rows were added later via translate-odm-parts / hand-drafting and
# don't exist in any workbook at all. Left in SHEETS above so it's still read and
# scanned for the same cached-error check as everything else (still useful to know
# about), but excluded here so a future run of this script can never silently
# overwrite real translation work with a stale, much smaller xlsx-derived copy.
NEVER_WRITE = {"translations"}

EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#SPILL!", "#CALC!"}


def strip_ws(v):
    return WS_RE.sub("", v) if isinstance(v, str) else v


def cell_to_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def read_sheet(wb_values, wb_formulas, sheet_name, pk_name):
    """Returns (header, rows, error_report) where error_report lists every row
    whose primary-key cell held a cached Excel error rather than a real value
    or a genuine blank -- these are NOT silently dropped without being surfaced."""
    ws_v = wb_values[sheet_name]
    ws_f = wb_formulas[sheet_name]

    raw_header = [c.value for c in ws_v[1]]
    real_cols = [i for i, h in enumerate(raw_header) if h not in (None, "")]
    header = [raw_header[i] for i in real_cols]
    pk_col_idx_in_header = header.index(pk_name)

    rows = []
    error_report = []
    for row_num in range(2, ws_v.max_row + 1):
        raw_row = [ws_v.cell(row=row_num, column=i + 1).value for i in real_cols]
        pk_val = raw_row[pk_col_idx_in_header]

        if pk_val in EXCEL_ERRORS:
            formula_present = isinstance(
                ws_f.cell(row=row_num, column=real_cols[pk_col_idx_in_header] + 1).value, str
            ) and ws_f.cell(row=row_num, column=real_cols[pk_col_idx_in_header] + 1).value.startswith("=")
            error_report.append((sheet_name, row_num, pk_val, formula_present))
            continue  # do not include -- but it IS reported, not silently lost

        if pk_val in (None, ""):
            continue  # genuinely blank template row, no formula error involved

        rows.append([cell_to_str(v) for v in raw_row])

    return header, rows, error_report


def fix_version(val):
    return val.replace("3.1.0", TARGET_VERSION) if isinstance(val, str) else val


def write_csv(basename, header, rows):
    version_row = ["Version", TARGET_VERSION] + [""] * (len(header) - 2)
    unsuffixed = os.path.join(OUT_DIR, f"{basename}.csv")
    with open(unsuffixed, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(version_row)
        w.writerow(header)
        w.writerows(rows)

    suffixed = os.path.join(OUT_DIR, f"{basename}_v{TARGET_VERSION}.csv")
    with open(suffixed, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)


def main(xlsx_path):
    wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False)

    all_error_reports = []
    sheet_data = {}

    for sheet_name, (basename, pk) in SHEETS.items():
        header, rows, error_report = read_sheet(wb_values, wb_formulas, sheet_name, pk)
        all_error_reports.extend(error_report)
        sheet_data[sheet_name] = (basename, pk, header, rows)

    # table-specific column trims
    for sheet_name, drop_cols in [("parts", DE_DW_TRIMMED), ("zones", {"Column1"})]:
        basename, pk, header, rows = sheet_data[sheet_name]
        keep_idx = [i for i, h in enumerate(header) if h not in drop_cols]
        new_header = [header[i] for i in keep_idx]
        new_rows = [[r[i] for i in keep_idx] for r in rows]
        sheet_data[sheet_name] = (basename, pk, new_header, new_rows)

    # partID whitespace trim, propagated to every table with a partID column
    for sheet_name in sheet_data:
        basename, pk, header, rows = sheet_data[sheet_name]
        if "partID" in header:
            idx = header.index("partID")
            for r in rows:
                r[idx] = strip_ws(r[idx])

    # wideNames' *Name columns are a cached display label for the *Input partID next
    # to them, typed once by whoever authored the row -- nothing in Excel re-syncs it
    # when that part's label changes later (see CLAUDE.md, "wideNames.csv's *Name
    # columns are cached labels"). Confirmed this is a live, ongoing risk, not just
    # a stale-snapshot artifact: even the most current xlsx source still had one
    # genuine drift (`popServ`'s label was renamed from "Wastewater treatment plant
    # designed capacity" to "Population Served" with nothing updating wideNames'
    # cached measureName). Resync every *Name to its *Input's current label on every
    # import, rather than trusting whatever text was typed into the workbook.
    wn_basename, wn_pk, wn_header, wn_rows = sheet_data["wideNames"]
    p_basename, p_pk, p_header, p_rows = sheet_data["parts"]
    s_basename, s_pk, s_header, s_rows = sheet_data["sets"]
    label_by_partid = {r[p_header.index("partID")]: r[p_header.index("label")] for r in p_rows}
    label_by_setpartid = {r[s_header.index("partID")]: r[s_header.index("label")] for r in s_rows}
    wn_pairs = ["reportTable", "partType", "compartment", "specimen", "fraction",
                "measure", "method", "unit", "aggregation", "attribute"]
    resynced = 0
    for r in wn_rows:
        for base in wn_pairs:
            input_idx = wn_header.index(f"{base}Input")
            name_idx = wn_header.index(f"{base}Name")
            input_val = r[input_idx]
            if not input_val or input_val.strip() in ("", "0"):
                continue
            real_label = (label_by_setpartid if base == "fraction" else label_by_partid).get(input_val)
            if real_label is not None and real_label != r[name_idx]:
                r[name_idx] = real_label
                resynced += 1
    if resynced:
        print(f"\nresynced {resynced} stale wideNames *Name label(s) to their *Input's current label")

    # 3.1.0 -> target version, everywhere
    total_version_fixes = 0
    for sheet_name in sheet_data:
        basename, pk, header, rows = sheet_data[sheet_name]
        for r in rows:
            for i, v in enumerate(r):
                if "3.1.0" in v:
                    r[i] = fix_version(v)
                    total_version_fixes += 1

    for sheet_name, (basename, pk, header, rows) in sheet_data.items():
        if sheet_name in NEVER_WRITE:
            print(f"{sheet_name}: {len(header)} cols, {len(rows)} rows -> SKIPPED (never written by this script)")
            continue
        write_csv(basename, header, rows)
        print(f"{sheet_name}: {len(header)} cols, {len(rows)} rows -> {basename}.csv")

    print(f"\n{total_version_fixes} cells corrected 3.1.0 -> {TARGET_VERSION}")

    if all_error_reports:
        print(f"\n{'='*70}")
        print(f"WARNING: {len(all_error_reports)} row(s) excluded due to a cached Excel")
        print("error on their primary-key cell -- NOT the same as a genuinely blank")
        print("row. This means a formula dependency was broken at the time this xlsx")
        print("was last saved. Resolve in the source workbook (fix the dependency,")
        print("force a recalculation, save) and re-run this script, or recover the")
        print("row's real values from a later/known-good snapshot if one exists.")
        print(f"{'='*70}")
        for sheet_name, row_num, err, has_formula in all_error_reports:
            print(f"  {sheet_name}!row{row_num}: primary key = {err}"
                  f"{' (formula present)' if has_formula else ' (no formula -- unexpected)'}")
        sys.exit(1)
    else:
        print("\nNo cached-error rows found on any primary-key column.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("usage: python3 import_from_xlsx.py <path-to-workingDoc.xlsx>")
        sys.exit(1)
    main(sys.argv[1])
